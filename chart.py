import openpyxl
from openpyxl. chart import LineChart, Reference
from datetime import datetime

x = datetime.now()
month = x.strftime("%b")

excelfilename = 'Chinese Futures Prices ZCM.xlsx'
path = r'C:\Users\asus\Desktop\KCH\Automation\ChineseFutures\%s' %excelfilename
excelgraph = 'Chinese Futures Prices ZCM.xlsx'
graphpath = r'C:\Users\asus\Desktop\KCH\Automation\ChineseFutures\%s' %excelgraph

wb = openpyxl.load_workbook(path)
ws = wb.active

maxrow = ws.max_row

print(maxrow)

def createchart():
    for i in range(3, maxrow):
        values = Reference(ws, min_col = 4, min_row = 4, max_col = 4, max_row = i)

    chart = LineChart()

    chart.add_data(values)

    chart.title = "Chinese Futures Price %s" %(month)

    chart.x_axis.title = "Time"

    chart.y_axis.title = "Price (yuan/t)"

    #chart.series[0].smooth = True

    ws.add_chart(chart, "G2")

    wb.save(graphpath)


if maxrow <= 5:
    createchart()
else:
    del ws._charts[0]
    createchart()